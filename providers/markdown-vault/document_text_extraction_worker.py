"""Resource-limited local text extraction worker for Markdown Vault documents."""

from __future__ import annotations

from collections.abc import Iterable
import json
import math
from pathlib import Path
import re
import resource
import shutil
import subprocess
import sys
import tempfile
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


class ExtractionFailure(RuntimeError):
    """Represent a bounded document failure returned to the provider process."""

    def __init__(
        self,
        detail: str,
        *,
        status_code: int = 422,
        code: str = "attachment_unextractable",
    ) -> None:
        """Create a structured extraction failure with a public-safe message."""
        super().__init__(detail)
        self.status_code = status_code
        self.code = code
        self.detail = detail


def _apply_resource_limits(timeout_seconds: float, max_result_bytes: int) -> None:
    """Limit CPU, address space, and output size before parsing untrusted input."""
    cpu_seconds = max(math.ceil(timeout_seconds), 1)
    resource.setrlimit(resource.RLIMIT_CPU, (cpu_seconds, cpu_seconds + 1))
    if sys.platform.startswith("linux"):
        memory_bytes = 768 * 1024 * 1024
        resource.setrlimit(resource.RLIMIT_AS, (memory_bytes, memory_bytes))
    resource.setrlimit(resource.RLIMIT_FSIZE, (max_result_bytes, max_result_bytes))


def _escape_cell(value: Any) -> str:
    """Normalize one Office table cell for a Markdown table."""
    return (
        str(value if value is not None else "")
        .strip()
        .replace("|", "\\|")
        .replace("\n", "<br>")
    )


def _markdown_table(rows: Iterable[Iterable[Any]]) -> str:
    """Convert rectangular or ragged Office rows into one Markdown table."""
    normalized = [[_escape_cell(value) for value in row] for row in rows]
    normalized = [row for row in normalized if any(row)]
    if not normalized:
        return ""
    width = max(len(row) for row in normalized)
    padded = [row + [""] * (width - len(row)) for row in normalized]
    return "\n".join(
        [
            "| " + " | ".join(padded[0]) + " |",
            "| " + " | ".join("---" for _ in range(width)) + " |",
            *("| " + " | ".join(row) + " |" for row in padded[1:]),
        ]
    )


def _bounded_markdown(parts: list[str], max_characters: int) -> str:
    """Join extracted blocks and enforce the provider's Markdown size contract."""
    markdown = "\n\n".join(part.strip() for part in parts if part.strip()).strip()
    if not markdown:
        raise ExtractionFailure("Document attachment contains no extractable text")
    if len(markdown) > max_characters:
        raise ExtractionFailure(
            "Converted attachment Markdown exceeds the configured size limit",
            status_code=413,
            code="attachment_markdown_too_large",
        )
    return markdown + "\n"


def _require_unit_limit(document_type: str, unit_count: int, max_units: int, unit_name: str) -> None:
    """Reject documents whose pages, slides, or sheets exceed the configured limit."""
    if unit_count > max_units:
        raise ExtractionFailure(
            f"{document_type} attachment exceeds the {max_units}-{unit_name} limit",
            status_code=413,
            code="attachment_page_limit_exceeded",
        )


def _extract_pdf(path: Path, max_units: int, max_characters: int) -> str:
    """Extract only the existing text layer from a PDF."""
    reader = PdfReader(str(path))
    if not reader.pages:
        raise ExtractionFailure("PDF attachment contains no pages")
    _require_unit_limit("PDF", len(reader.pages), max_units, "page")
    parts: list[str] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = str(page.extract_text() or "").strip()
        if text:
            parts.append(f"## Page {page_number}\n\n{text}")
    if not parts:
        raise ExtractionFailure("PDF attachment contains no extractable text")
    return _bounded_markdown(parts, max_characters)


def _docx_paragraph_markdown(paragraph: Paragraph) -> str:
    """Preserve basic heading and list semantics for one Word paragraph."""
    text = paragraph.text.strip()
    if not text:
        return ""
    style_name = str(paragraph.style.name if paragraph.style is not None else "")
    heading = re.fullmatch(r"Heading ([1-6])", style_name)
    if heading:
        return f"{'#' * int(heading.group(1))} {text}"
    if style_name.startswith("List"):
        return f"- {text}"
    return text


def _extract_docx(path: Path, max_characters: int) -> str:
    """Extract Word paragraphs and tables in document order."""
    document = Document(str(path))
    parts: list[str] = []
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            block = _docx_paragraph_markdown(Paragraph(child, document))
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            block = _markdown_table([[cell.text for cell in row.cells] for row in table.rows])
        else:
            continue
        if block:
            parts.append(block)
    return _bounded_markdown(parts, max_characters)


def _iter_pptx_shapes(shapes: Any) -> Iterable[Any]:
    """Yield PowerPoint shapes in visual order, flattening group containers."""
    for shape in sorted(shapes, key=lambda item: (int(item.top), int(item.left))):
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_pptx_shapes(shape.shapes)
        else:
            yield shape


def _extract_pptx(path: Path, max_units: int, max_characters: int) -> str:
    """Extract slide text and tables without rendering presentation graphics."""
    presentation = Presentation(str(path))
    _require_unit_limit("PowerPoint", len(presentation.slides), max_units, "slide")
    parts: list[str] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_parts: list[str] = []
        for shape in _iter_pptx_shapes(slide.shapes):
            if bool(getattr(shape, "has_table", False)):
                table = _markdown_table(
                    [[cell.text for cell in row.cells] for row in shape.table.rows]
                )
                if table:
                    slide_parts.append(table)
                continue
            text = str(getattr(shape, "text", "") or "").strip()
            if text:
                slide_parts.append(text)
        if slide_parts:
            parts.append(f"## Slide {slide_number}\n\n" + "\n\n".join(slide_parts))
    return _bounded_markdown(parts, max_characters)


def _extract_xlsx(path: Path, max_units: int, max_characters: int) -> str:
    """Extract spreadsheet cell values by sheet and row."""
    workbook = load_workbook(
        filename=str(path),
        read_only=True,
        data_only=True,
        keep_links=False,
    )
    try:
        _require_unit_limit("Excel", len(workbook.sheetnames), max_units, "sheet")
        parts: list[str] = []
        for worksheet in workbook.worksheets:
            rows = [
                list(row)
                for row in worksheet.iter_rows(values_only=True)
                if any(value is not None for value in row)
            ]
            table = _markdown_table(rows)
            if table:
                parts.append(f"## Sheet: {worksheet.title}\n\n{table}")
        return _bounded_markdown(parts, max_characters)
    finally:
        workbook.close()


def _convert_legacy_office(path: Path, timeout_seconds: float) -> Path:
    """Convert one legacy Office binary when the optional LibreOffice command is installed."""
    libreoffice = shutil.which("libreoffice") or shutil.which("soffice")
    if libreoffice is None:
        raise ExtractionFailure(
            "LibreOffice is required to extract text from .doc, .ppt, and .xls attachments",
            status_code=503,
            code="office_converter_unavailable",
        )
    target_extension = {".doc": ".docx", ".ppt": ".pptx", ".xls": ".xlsx"}[path.suffix.lower()]
    output_dir = path.parent / "converted"
    output_dir.mkdir()
    with tempfile.TemporaryDirectory(prefix="libreoffice-profile-", dir=path.parent) as profile_dir:
        try:
            completed = subprocess.run(
                [
                    libreoffice,
                    "--headless",
                    "--nologo",
                    "--nodefault",
                    "--nofirststartwizard",
                    f"-env:UserInstallation={Path(profile_dir).as_uri()}",
                    "--convert-to",
                    target_extension.removeprefix("."),
                    "--outdir",
                    str(output_dir),
                    str(path),
                ],
                capture_output=True,
                text=True,
                timeout=max(timeout_seconds, 1),
                check=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise ExtractionFailure(
                f"LibreOffice conversion exceeded the request deadline: {path.name}",
                status_code=504,
                code="attachment_conversion_timeout",
            ) from exc
    converted_path = output_dir / f"{path.stem}{target_extension}"
    if completed.returncode != 0 or not converted_path.is_file():
        raise ExtractionFailure(f"LibreOffice could not convert attachment: {path.name}")
    return converted_path


def _extract(
    path: Path,
    max_units: int,
    max_characters: int,
    timeout_seconds: float,
) -> dict[str, Any]:
    """Dispatch a supported document container to its local text extractor."""
    extension = path.suffix.lower()
    if extension in {".doc", ".ppt", ".xls"}:
        path = _convert_legacy_office(path, timeout_seconds)
        extension = path.suffix.lower()
    extractors = {
        ".pdf": lambda: _extract_pdf(path, max_units, max_characters),
        ".docx": lambda: _extract_docx(path, max_characters),
        ".pptx": lambda: _extract_pptx(path, max_units, max_characters),
        ".xlsx": lambda: _extract_xlsx(path, max_units, max_characters),
    }
    extractor = extractors.get(extension)
    if extractor is None:
        raise ExtractionFailure(f"Unsupported document type for local extraction: {extension}")
    return {"markdown": extractor()}


def main(argv: list[str] | None = None) -> int:
    """Run one isolated extraction and write a structured result for the parent process."""
    arguments = list(sys.argv[1:] if argv is None else argv)
    if len(arguments) != 6:
        return 2
    input_path = Path(arguments[0])
    result_path = Path(arguments[1])
    try:
        max_units = int(arguments[2])
        max_characters = int(arguments[3])
        timeout_seconds = float(arguments[4])
        max_result_bytes = int(arguments[5])
        _apply_resource_limits(timeout_seconds, max_result_bytes)
        payload = _extract(input_path, max_units, max_characters, timeout_seconds)
    except ExtractionFailure as exc:
        payload = {
            "error": {
                "statusCode": exc.status_code,
                "code": exc.code,
                "detail": exc.detail,
            }
        }
    except Exception as exc:
        payload = {
            "error": {
                "statusCode": 422,
                "code": "attachment_unextractable",
                "detail": f"Local document extraction failed: {type(exc).__name__}",
            }
        }
    result_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
