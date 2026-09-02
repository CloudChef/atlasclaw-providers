"""Critical local-text and visual-image tests for Knowledge attachments."""

from __future__ import annotations

import importlib.util
from io import BytesIO
from pathlib import Path
import sys
from types import ModuleType
from typing import Any

from docx import Document
from openpyxl import Workbook
import pytest
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from pptx import Presentation
from pptx.util import Inches


PROVIDER_ROOT = Path(__file__).resolve().parents[1]


def _load_converter_module() -> ModuleType:
    """Load the provider converter without relying on an installed package name."""
    converter_path = PROVIDER_ROOT / "knowledge_attachment_converter.py"
    spec = importlib.util.spec_from_file_location(
        "markdown_vault_knowledge_attachment_converter_test",
        converter_path,
    )
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


converter_module = _load_converter_module()


class _FakeVisualRuntime:
    """Capture provider visual calls through the same narrow contract Core exposes."""

    def __init__(self, *, model_name: str = "vision-test", available: bool = True) -> None:
        """Create a deterministic visual bridge for one converter test."""
        self.visual_model_name = model_name
        self.visual_available = available
        self.calls: list[dict[str, Any]] = []

    async def run_visual(self, **kwargs: Any) -> str:
        """Record one visual request and return image-shaped Markdown."""
        self.calls.append(kwargs)
        return "IMAGE_MARKDOWN_MARKER_8834"


def _text_pdf(marker: str) -> bytes:
    """Build a one-page PDF with a real extractable text layer."""
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    content = DecodedStreamObject()
    content.set_data(f"BT /F1 12 Tf 72 720 Td ({marker}) Tj ET".encode())
    page[NameObject("/Contents")] = content
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def _blank_pdf(page_count: int = 1) -> bytes:
    """Build a PDF without an extractable text layer."""
    writer = PdfWriter()
    for _ in range(page_count):
        writer.add_blank_page(width=612, height=792)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def _docx(marker: str) -> bytes:
    """Build a Word document containing a heading, paragraph, and table."""
    document = Document()
    document.add_heading("Word title", level=1)
    document.add_paragraph(marker)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Key"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Certification"
    table.cell(1, 1).text = "ISO 27001"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _pptx(marker: str) -> bytes:
    """Build a PowerPoint document containing searchable slide text."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Slide title"
    textbox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1))
    textbox.text = marker
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def _xlsx(marker: str) -> bytes:
    """Build an Excel workbook containing a bounded table."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Assessment"
    worksheet.append(["Organization", "Result"])
    worksheet.append([marker, "Leader"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


async def _convert_document(
    file_name: str, content_type: str, data: bytes
) -> tuple[dict[str, Any], _FakeVisualRuntime]:
    """Convert one local-text document through the public converter entrypoint."""
    runtime = _FakeVisualRuntime()
    converter = converter_module.KnowledgeAttachmentLlmConverter(runtime, timeout_seconds=30)
    result = await converter.convert(
        {
            "fileId": "document-1",
            "fileName": file_name,
            "contentType": content_type,
            "data": data,
        },
        {"knowledgeId": "knowledge-1"},
    )
    return result, runtime


@pytest.mark.asyncio
async def test_pdf_text_layer_is_extracted_locally_without_visual_model() -> None:
    """Verify PDF extraction uses only the existing text layer."""
    result, runtime = await _convert_document(
        "evidence.pdf", "application/pdf", _text_pdf("PDF_MARKER_1042")
    )

    assert "## Page 1" in result["markdown"]
    assert "PDF_MARKER_1042" in result["markdown"]
    assert result["assets"] == []
    assert runtime.calls == []


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("file_name", "content_type", "data", "marker"),
    [
        (
            "guide.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            _docx("DOCX_MARKER_2043"),
            "DOCX_MARKER_2043",
        ),
        (
            "slides.pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            _pptx("PPTX_MARKER_3044"),
            "PPTX_MARKER_3044",
        ),
        (
            "assessment.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            _xlsx("XLSX_MARKER_4045"),
            "XLSX_MARKER_4045",
        ),
    ],
)
async def test_office_text_is_extracted_locally(
    file_name: str, content_type: str, data: bytes, marker: str
) -> None:
    """Verify supported Office containers become Markdown without rendering or a model."""
    result, runtime = await _convert_document(file_name, content_type, data)

    assert marker in result["markdown"]
    assert result["assets"] == []
    assert runtime.calls == []


@pytest.mark.asyncio
async def test_image_is_delegated_to_visual_model() -> None:
    """Verify a direct image attachment uses Core's visual-model bridge."""
    runtime = _FakeVisualRuntime()
    converter = converter_module.KnowledgeAttachmentLlmConverter(runtime, timeout_seconds=30)

    result = await converter.convert(
        {
            "fileId": "image-1",
            "fileName": "diagram.png",
            "contentType": "image/png",
            "data": b"\x89PNG\r\n\x1a\nimage-data",
        },
        {"knowledgeId": "knowledge-1"},
    )

    assert "## Image" in result["markdown"]
    assert "IMAGE_MARKDOWN_MARKER_8834" in result["markdown"]
    assert result["assets"][0]["fileName"] == "image.png"
    assert runtime.calls[0]["media_type"] == "image/png"


@pytest.mark.asyncio
async def test_textless_pdf_is_rejected_without_visual_processing() -> None:
    """Verify scanned or blank PDFs do not silently enter image-model processing."""
    runtime = _FakeVisualRuntime()
    converter = converter_module.KnowledgeAttachmentLlmConverter(runtime, timeout_seconds=30)

    with pytest.raises(
        converter_module.KnowledgeAttachmentConversionError,
        match="PDF attachment contains no extractable text",
    ):
        await converter.convert(
            {
                "fileId": "scan-1",
                "fileName": "scan.pdf",
                "contentType": "application/pdf",
                "data": _blank_pdf(),
            },
            {"knowledgeId": "knowledge-1"},
        )

    assert runtime.calls == []


@pytest.mark.asyncio
async def test_text_only_deepseek_rejects_direct_image() -> None:
    """Verify a text-only DeepSeek model is not used for image attachments."""
    runtime = _FakeVisualRuntime(model_name="deepseek-chat")
    converter = converter_module.KnowledgeAttachmentLlmConverter(runtime, timeout_seconds=30)

    with pytest.raises(
        converter_module.KnowledgeAttachmentConversionError,
        match="does not support visual input",
    ):
        await converter.convert(
            {
                "fileId": "image-1",
                "fileName": "diagram.png",
                "contentType": "image/png",
                "data": b"\x89PNG\r\n\x1a\nimage-data",
            },
            {"knowledgeId": "knowledge-1"},
        )

    assert runtime.calls == []


@pytest.mark.asyncio
async def test_pdf_page_limit_is_enforced_before_extraction() -> None:
    """Verify PDF page limits remain enforced without rendering or model calls."""
    runtime = _FakeVisualRuntime()
    converter = converter_module.KnowledgeAttachmentLlmConverter(
        runtime,
        timeout_seconds=30,
        max_document_pages=1,
    )

    with pytest.raises(converter_module.KnowledgeAttachmentConversionError) as error:
        await converter.convert(
            {
                "fileId": "pdf-many",
                "fileName": "many-pages.pdf",
                "contentType": "application/pdf",
                "data": _blank_pdf(2),
            },
            {"knowledgeId": "knowledge-1"},
        )

    assert error.value.status_code == 413
    assert error.value.code == "attachment_page_limit_exceeded"
    assert runtime.calls == []


@pytest.mark.asyncio
async def test_legacy_office_format_remains_unsupported() -> None:
    """Verify legacy OLE Office formats stay outside the local extraction boundary."""
    converter = converter_module.KnowledgeAttachmentLlmConverter(
        _FakeVisualRuntime(available=False),
        timeout_seconds=30,
    )

    with pytest.raises(
        converter_module.KnowledgeAttachmentConversionError,
        match="Unsupported attachment type",
    ):
        await converter.convert(
            {
                "fileId": "doc-legacy",
                "fileName": "legacy.doc",
                "contentType": "application/msword",
                "data": b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1legacy-office",
            },
            {"knowledgeId": "knowledge-1"},
        )
